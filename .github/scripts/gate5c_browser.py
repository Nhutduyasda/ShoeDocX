import hashlib
import json
import os
from pathlib import Path
import subprocess
import tempfile
import time
import urllib.error
import urllib.request
import uuid
import openpyxl


def browser(ci):
    feature = ci.ROOT / 'feature'
    folder = ci.EVIDENCE / 'gate5c'; folder.mkdir(exist_ok=True)
    environment = os.environ.copy()
    environment.update(GATE5C_DATABASE=os.environ['GATE5C_SQL'] + ';Database=gate5c_browser_' + uuid.uuid4().hex,
                       STUDIO_DEMO_PORT='5289', XNK_API_TARGET='http://127.0.0.1:5289', STUDIO_BASE_URL='http://127.0.0.1:5177', STUDIO_EVIDENCE_DIR=str(folder))
    host = str(feature / 'tests/Gate5CDemoHost/bin/Debug/net8.0/Gate5CDemoHost.dll')
    def command(args):
        return subprocess.run(['dotnet', host, *args], cwd=feature, env=environment, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=90)
    assert command(['--migrate']).returncode == 0, 'Ephemeral browser SQL migration failed'
    info = command(['--sql-info']); assert info.returncode == 0
    (folder / 'sql-server.json').write_text(json.dumps(json.loads(info.stdout), indent=2))
    processes = []; logs = []
    with tempfile.TemporaryDirectory(prefix='gate5c-private-logs-') as temporary:
        def launch(args, cwd):
            log = open(Path(temporary) / (str(len(logs)) + '.log'), 'wb'); logs.append(log)
            process = subprocess.Popen(args, cwd=cwd, env=environment, stdout=log, stderr=subprocess.STDOUT); processes.append(process); return process
        def ready():
            deadline = time.monotonic() + 90
            for url in ('http://127.0.0.1:5289/api/universal-doc/bootstrap', 'http://127.0.0.1:5177/universal-doc-studio'):
                while True:
                    try:
                        with urllib.request.urlopen(url, timeout=2): break
                    except urllib.error.HTTPError as e:
                        assert e.code == 401, 'Readiness HTTP error'; break
                    except (urllib.error.URLError, TimeoutError):
                        assert time.monotonic() < deadline, 'Browser host startup timed out'; time.sleep(.2)
        try:
            backend = launch(['dotnet', host], feature)
            launch(['node', str(feature / 'frontend/node_modules/vite/bin/vite.js'), '--host', '127.0.0.1', '--port', '5177', '--strictPort'], feature / 'frontend')
            ready()
            output = open(Path(temporary) / 'browser.log', 'wb'); logs.append(output)
            browser = subprocess.Popen(['node', 'frontend/tests/universalDocGate5C.e2e.mjs'], cwd=feature, env=environment, stdout=output, stderr=subprocess.STDOUT); processes.append(browser)
            deadline = time.monotonic() + 240; restarted = False
            while browser.poll() is None:
                assert time.monotonic() < deadline, 'Gate 5C browser timed out'
                if not restarted and (folder / 'restart-requested').exists():
                    backend.terminate(); backend.wait(timeout=15)
                    backend = launch(['dotnet', host], feature); ready()
                    (folder / 'restart-completed').write_text('actual backend process restarted')
                    restarted = True
                time.sleep(.2)
            if browser.returncode:
                output.flush()
                print('\n'.join(line for line in (Path(temporary) / 'browser.log').read_text(encoding='utf-8').splitlines() if any(t in line for t in ('Error:', 'waiting for', 'locator.', 'Expected:', 'Actual:'))), flush=True)
            assert browser.returncode == 0 and restarted, 'Real restarted browser acceptance failed'
            summary = json.loads((folder / 'gate5c-summary.json').read_text())
            assert summary['status'] == 'PASS' and len(summary['checks']) == 16 and not summary['runtimeErrors']
            path = folder / 'synthetic-gate5c.xlsx'; workbook = openpyxl.load_workbook(path)
            assert workbook.sheetnames == ['Data1', 'Data2', 'Reference']
            for name in ['Data1', 'Data2']:
                ws = workbook[name]; assert ws['B1'].value == 15
                for row, identifier in [(4, '0001'), (5, '0002')]:
                    assert [ws.cell(row, c).value for c in range(1, 5)] == [identifier, 2.5, 3, 7.5]
                    assert ws.cell(row, 5).value == f'=B{row}*C{row}'
                    assert ws.cell(row, 2).number_format == '0.00' and ws.cell(row, 2).border.bottom.style == 'thin'
                assert ws['E6'].value == '=SUM(E4:E5)' and 'A7:C7' in [str(x) for x in ws.merged_cells.ranges]
            assert workbook['Reference']['A1'].value == 'Dữ liệu tham chiếu 00123'
            data = path.read_bytes()
            (folder / 'workbook-qa.json').write_text(json.dumps({'status': 'PASS', 'independent_openpyxl': True, 'bytes': len(data), 'sha256': hashlib.sha256(data).hexdigest(), 'sheets': 2, 'values_formats_formulas_merges_reference': True}, indent=2))
            path.unlink()
            print('Gate 5C: 16 actual authenticated browser checks + real process restart + independent downloaded XLSX reopen PASS', flush=True)
        finally:
            for p in processes:
                if p.poll() is None: p.terminate()
            for p in processes:
                try: p.wait(timeout=10)
                except subprocess.TimeoutExpired: p.kill(); p.wait()
            for f in logs: f.close()
            assert command(['--drop']).returncode == 0, 'Ephemeral browser database cleanup failed'
