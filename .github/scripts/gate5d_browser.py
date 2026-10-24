import hashlib
import json
import os
from pathlib import Path
import subprocess
import tempfile
import time
import urllib.error
import urllib.request
import uuid
import openpyxl


def browser(ci):
    feature = ci.ROOT / 'feature'
    folder = ci.EVIDENCE / 'gate5d'; folder.mkdir(exist_ok=True)
    private = tempfile.TemporaryDirectory(prefix='gate5d-synthetic-private-')
    private_folder = Path(private.name)
    environment = os.environ.copy()
    environment.update(GATE5C_DATABASE=os.environ['GATE5C_SQL'] + ';Database=gate5c_browser_' + uuid.uuid4().hex,
                       STUDIO_DEMO_PORT='5291', XNK_API_TARGET='http://127.0.0.1:5291', STUDIO_BASE_URL='http://127.0.0.1:5179', STUDIO_EVIDENCE_DIR=str(private_folder))
    host = str(feature / 'tests/Gate5DDemoHost/bin/Debug/net8.0/Gate5DDemoHost.dll')
    def command(args):
        return subprocess.run(['dotnet', host, *args], cwd=feature, env=environment, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=90)
    assert command(['--migrate']).returncode == 0, 'Ephemeral browser SQL migration failed'
    processes = []; logs = []
    with tempfile.TemporaryDirectory(prefix='gate5d-private-logs-') as temporary:
        def launch(args, cwd):
            log = open(Path(temporary) / (str(len(logs)) + '.log'), 'wb'); logs.append(log)
            process = subprocess.Popen(args, cwd=cwd, env=environment, stdout=log, stderr=subprocess.STDOUT); processes.append(process); return process
        def ready():
            deadline = time.monotonic() + 90
            for url in ('http://127.0.0.1:5291/api/universal-doc/bootstrap', 'http://127.0.0.1:5179/universal-doc-studio'):
                while True:
                    try:
                        with urllib.request.urlopen(url, timeout=2): break
                    except urllib.error.HTTPError as e:
                        assert e.code == 401, 'Readiness HTTP error'; break
                    except (urllib.error.URLError, TimeoutError):
                        assert time.monotonic() < deadline, 'Browser host startup timed out'; time.sleep(.2)
        try:
            backend = launch(['dotnet', host], feature)
            launch(['node', str(feature / 'frontend/node_modules/vite/bin/vite.js'), '--host', '127.0.0.1', '--port', '5179', '--strictPort'], feature / 'frontend')
            ready()
            output = open(Path(temporary) / 'browser.log', 'wb'); logs.append(output)
            browser = subprocess.Popen(['node', 'frontend/tests/universalDocGate5D.e2e.mjs'], cwd=feature, env=environment, stdout=output, stderr=subprocess.STDOUT); processes.append(browser)
            deadline = time.monotonic() + 720; restarted = False
            while browser.poll() is None:
                assert time.monotonic() < deadline, 'Gate 5C browser timed out'
                if not restarted and (private_folder / 'restart-requested').exists():
                    backend.terminate(); backend.wait(timeout=15)
                    backend = launch(['dotnet', host], feature); ready()
                    (private_folder / 'restart-completed').write_text('actual backend process restarted')
                    restarted = True
                time.sleep(.2)
            if browser.returncode:
                output.flush()
                print('\n'.join(line for line in (Path(temporary) / 'browser.log').read_text(encoding='utf-8').splitlines() if any(t in line for t in ('Error:', 'waiting for', 'locator.', 'Expected:', 'Actual:'))), flush=True)
            assert browser.returncode == 0 and restarted, 'Real restarted browser acceptance failed'
            summary = json.loads((private_folder / 'office-summary.json').read_text(encoding='utf-8'))
            assert summary['status'] == 'PASS' and len(set(summary['journeys'])) == 5 and not summary['runtimeErrors']
            (folder / 'office-summary.json').write_text(json.dumps(summary, indent=2), encoding='utf-8')
            downloads = json.loads((private_folder / 'office-downloads.json').read_text(encoding='utf-8'))
            assert len(downloads) == 4
            qa = [verify_workbook(private_folder / d['name'], d['scenario'], d['rows']) for d in downloads]
            corpus_folder = private_folder / 'corpus'; result = command(['--corpus', str(corpus_folder)])
            if result.returncode: print('Corpus process failed: ' + str(result.returncode), flush=True)
            assert result.returncode == 0, 'Actual corpus pipeline failed'
            manifest = json.loads((corpus_folder / 'corpus-manifest.json').read_text(encoding='utf-8'))
            supported = [m for m in manifest if m['status'] == 'SUPPORTED_EXPORTED']
            rejected = [m for m in manifest if m['status'] == 'REJECTED']
            assert {m['scenario'] for m in supported} == set('ABCDEFG')
            assert len(rejected) == 4 and all(not m['outputBytesReleased'] for m in rejected)
            qa += [verify_workbook(corpus_folder / (m['scenario'] + '.xlsx'), m['scenario'], m['rows']) for m in supported]
            (folder / 'workbook-qa.json').write_text(json.dumps(qa, indent=2), encoding='utf-8')
            (folder / 'compatibility.json').write_text(json.dumps(manifest, indent=2), encoding='utf-8')
            print('Gate 5D: five actual authenticated office journeys + actual SQL restart + 11 independent XLSX reopen checks PASS', flush=True)
        finally:
            for p in processes:
                if p.poll() is None: p.terminate()
            for p in processes:
                try: p.wait(timeout=10)
                except subprocess.TimeoutExpired: p.kill(); p.wait()
            for f in logs: f.close()
            assert command(['--drop']).returncode == 0, 'Ephemeral browser database cleanup failed'
            private.cleanup()


def verify_workbook(path, scenario, rows):
    import datetime
    workbook = openpyxl.load_workbook(path)
    names = {'B':['NhapKho'], 'C':['XuatKho'], 'D':['Packing'], 'E':['Kho','TongHop'], 'F':['Kho','TongHop','ChiTiet']}.get(scenario, ['HoaDon'])
    assert workbook.sheetnames == names + ['ThamChieu']
    for name in names:
        ws = workbook[name]
        assert ws['B1'].value == 'CT-00123' and ws['B2'].value == datetime.datetime(2026,10,24) and ws['B3'].value is True
        assert ws['B2'].number_format == 'yyyy-mm-dd'
        for row in range(6, 6 + rows):
            assert [ws.cell(row,c).value for c in range(1,6)] == ['0012'+str(row-5),2.5,3,7.5,'cái']
            assert ws.cell(row,1).data_type == 's' and ws.cell(row,6).value == f'=B{row}*C{row}'
            assert ws.cell(row,2).number_format == '0.00' and ws.cell(row,2).border.bottom.style == 'thin'
        assert ws.cell(6+rows,6).value == f'=SUM(F6:F{5+rows})'
        assert f'A{7+rows}:D{7+rows}' in [str(r) for r in ws.merged_cells.ranges]
    assert workbook['ThamChieu']['A1'].value == 'Tham chiếu giữ nguyên 00123'
    data = path.read_bytes()
    return {'scenario':scenario, 'rows':rows, 'mapped_sheets':len(names), 'status':'PASS', 'independent_openpyxl':True, 'bytes':len(data), 'sha256':hashlib.sha256(data).hexdigest(), 'values_types_formats_formulas_borders_merges_references':True}
